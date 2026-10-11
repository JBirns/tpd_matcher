#!/usr/bin/env python3
"""Allocate doctors to hospitals from the ranking-sheet format described in README."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


class MatchingError(ValueError):
    """Raised when input data cannot describe a valid matching."""


UNALLOCATED_HOSPITAL = "UNALLOCATED"


@dataclass(frozen=True)
class Doctor:
    name: str
    employment: str  # "full-time" or "part-time"
    rankings: tuple[str, ...]
    row_number: int

    @property
    def is_part_time(self) -> bool:
        return self.employment == "part-time"


@dataclass(frozen=True)
class Hospital:
    name: str
    places: int
    allows_jobshares: bool


def normalise_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def parse_employment(value: object, row_number: int) -> str:
    value = normalise_text(value).casefold().replace("_", "-")
    if value in {"full-time", "full time", "fulltime", "ft", "f", "full"}:
        return "full-time"
    if value in {"part-time", "part time", "parttime", "pt", "p", "part"}:
        return "part-time"
    raise MatchingError(
        f"Row {row_number}: column D must say full-time or part-time (got {value!r})."
    )


def read_rows(path: Path) -> Iterable[list[object]]:
    if path.suffix.casefold() == ".csv":
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            yield from csv.reader(handle)
        return
    if path.suffix.casefold() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise MatchingError("Reading .xlsx requires openpyxl; run: python -m pip install -r requirements.txt") from error
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            yield from workbook.active.iter_rows(values_only=True)
        finally:
            workbook.close()
        return
    raise MatchingError("Ranking file must be a .csv or .xlsx file.")


def read_doctors(path: Path, skip_header: bool) -> list[Doctor]:
    doctors: list[Doctor] = []
    seen_names: set[str] = set()
    for row_number, row in enumerate(read_rows(path), start=1):
        if skip_header and row_number == 1:
            continue
        values = list(row) + [""] * max(0, 6 - len(row))
        name = normalise_text(values[2])
        # Empty rows are harmless, including trailing spreadsheet rows.
        if not name:
            continue
        if name in seen_names:
            raise MatchingError(f"Row {row_number}: duplicate doctor name {name!r}.")
        seen_names.add(name)
        rankings = tuple(
            hospital for hospital in (normalise_text(item) for item in values[5:]) if hospital
        )
        if len(rankings) != len(set(rankings)):
            raise MatchingError(f"Row {row_number}: {name!r} ranks a hospital more than once.")
        doctors.append(Doctor(name, parse_employment(values[3], row_number), rankings, row_number))
    if not doctors:
        raise MatchingError("No doctors were found in column C.")
    return doctors


def read_configuration(path: Path) -> tuple[dict[str, Hospital], dict[str, str]]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise MatchingError(f"Could not read configuration JSON: {error}") from error
    raw_hospitals = data.get("hospitals")
    if not isinstance(raw_hospitals, dict) or not raw_hospitals:
        raise MatchingError("Configuration needs a non-empty 'hospitals' object.")
    hospitals: dict[str, Hospital] = {}
    for name, setting in raw_hospitals.items():
        name = normalise_text(name)
        if not name or not isinstance(setting, dict):
            raise MatchingError("Each hospital needs a non-empty name and an object of settings.")
        places = setting.get("full_time_places")
        allows_jobshares = setting.get("allows_jobshares")
        if isinstance(places, bool) or not isinstance(places, int) or places < 0:
            raise MatchingError(f"Hospital {name!r}: full_time_places must be a non-negative integer.")
        if not isinstance(allows_jobshares, bool):
            raise MatchingError(f"Hospital {name!r}: allows_jobshares must be true or false.")
        hospitals[name] = Hospital(name, places, allows_jobshares)
    raw_preallocations = data.get("preallocations", {})
    if not isinstance(raw_preallocations, dict):
        raise MatchingError("'preallocations' must map doctor names to hospital names.")
    preallocations = {normalise_text(doctor): normalise_text(hospital)
                      for doctor, hospital in raw_preallocations.items()}
    return hospitals, preallocations


def weight(doctor: Doctor, hospital: Hospital) -> int:
    """Capacity in half-place units. PT uses a half-place only at jobshare hospitals."""
    return 1 if doctor.is_part_time and hospital.allows_jobshares else 2


def match(doctors: list[Doctor], hospitals: dict[str, Hospital], preallocations: dict[str, str]) -> tuple[dict[str, list[str]], list[str]]:
    if UNALLOCATED_HOSPITAL in hospitals:
        raise MatchingError(f"{UNALLOCATED_HOSPITAL!r} is reserved for doctors who cannot be matched.")
    doctor_by_name = {doctor.name: doctor for doctor in doctors}
    unknown_doctors = set(preallocations) - set(doctor_by_name)
    if unknown_doctors:
        raise MatchingError(f"Pre-allocation names not in spreadsheet: {', '.join(sorted(unknown_doctors))}.")
    unknown_hospitals = set(preallocations.values()) - set(hospitals)
    if unknown_hospitals:
        raise MatchingError(f"Pre-allocation hospitals not configured: {', '.join(sorted(unknown_hospitals))}.")
    for doctor in doctors:
        unknown = set(doctor.rankings) - set(hospitals)
        if unknown:
            raise MatchingError(f"{doctor.name!r} ranks unknown hospital(s): {', '.join(sorted(unknown))}.")

    remaining = {name: hospital.places * 2 for name, hospital in hospitals.items()}
    assignments: dict[str, list[str]] = defaultdict(list)
    assigned: set[str] = set()
    # Locks are placed first and do not need to appear in the doctor's ranking.
    for doctor in doctors:
        hospital_name = preallocations.get(doctor.name)
        if hospital_name is None:
            continue
        hospital = hospitals[hospital_name]
        required = weight(doctor, hospital)
        if remaining[hospital_name] < required:
            raise MatchingError(f"Pre-allocations exceed capacity at {hospital_name!r}.")
        remaining[hospital_name] -= required
        assignments[hospital_name].append(doctor.name)
        assigned.add(doctor.name)

    # Doctor-proposing deferred acceptance. Hospitals have no preference ordering in
    # the supplied data, so an accepted doctor is never displaced; spreadsheet row
    # order supplies the deterministic tie-break when capacity is contested.
    unmatched: list[str] = []
    for doctor in doctors:
        if doctor.name in assigned:
            continue
        for hospital_name in doctor.rankings:
            hospital = hospitals[hospital_name]
            required = weight(doctor, hospital)
            if remaining[hospital_name] >= required:
                remaining[hospital_name] -= required
                assignments[hospital_name].append(doctor.name)
                assigned.add(doctor.name)
                break
        else:
            unmatched.append(doctor.name)
    return assignments, unmatched


def write_output(path: Path, hospitals: dict[str, Hospital], assignments: dict[str, list[str]],
                 unmatched: list[str]) -> None:
    columns = list(hospitals)
    output_assignments = dict(assignments)
    if unmatched:
        columns.append(UNALLOCATED_HOSPITAL)
        output_assignments[UNALLOCATED_HOSPITAL] = unmatched
    height = max((len(output_assignments.get(name, [])) for name in columns), default=0)
    rows = [
        [output_assignments[name][index] if index < len(output_assignments.get(name, [])) else ""
         for name in columns]
        for index in range(height)
    ]
    if path.suffix.casefold() == ".csv":
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            writer.writerows(rows)
        return
    if path.suffix.casefold() == ".xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as error:
            raise MatchingError("Writing .xlsx requires openpyxl; run: python -m pip install -r requirements.txt") from error
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Assignments"
        worksheet.append(columns)
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        return
    raise MatchingError("Output file must have a .csv or .xlsx extension.")


def show_fullscreen_message(parent: object, title: str, message: str,
                            *, use_parent: bool = False, require_ok: bool = False) -> None:
    """Show long instructions as a full-screen page with wrapping, scrollable text."""
    import tkinter as tk
    from tkinter import ttk

    window = parent if use_parent else tk.Toplevel(parent)
    window.title(title)
    if not use_parent:
        window.transient(parent)
    window.attributes("-fullscreen", True)
    window.columnconfigure(0, weight=1)
    window.rowconfigure(0, weight=1)

    content = ttk.Frame(window, padding=32)
    content.grid(sticky="nsew")
    content.columnconfigure(0, weight=1)
    content.rowconfigure(1, weight=1)
    ttk.Label(content, text=title, font=("TkDefaultFont", 18, "bold")).grid(
        row=0, column=0, columnspan=2, pady=(0, 18), sticky="w"
    )
    text = tk.Text(content, wrap="word", relief="solid", borderwidth=1, padx=14, pady=14,
                   font=("TkDefaultFont", 12))
    scrollbar = ttk.Scrollbar(content, orient="vertical", command=text.yview)
    text.configure(yscrollcommand=scrollbar.set)
    text.insert("1.0", message)
    text.configure(state="disabled")
    text.grid(row=1, column=0, sticky="nsew")
    scrollbar.grid(row=1, column=1, sticky="ns")

    actions = ttk.Frame(content)
    actions.grid(row=2, column=0, columnspan=2, pady=(18, 0), sticky="e")
    done = tk.BooleanVar(master=window, value=False)

    def close() -> None:
        if use_parent:
            done.set(True)
        else:
            window.grab_release()
            window.destroy()

    ttk.Button(actions, text="OK", command=close).grid(row=0, column=0)
    window.protocol("WM_DELETE_WINDOW", (lambda: None) if require_ok else close)
    if not require_ok:
        window.bind("<Escape>", lambda _event: close())

    if use_parent:
        window.wait_variable(done)
        for child in window.winfo_children():
            child.destroy()
        window.attributes("-fullscreen", False)
    else:
        window.update_idletasks()
        window.lift()
        window.focus_force()
        window.grab_set()
        window.wait_window()


def show_error(message: str) -> None:
    """Make startup errors visible when running a windowed packaged application."""
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        messagebox.showerror("TPD Matcher", message, parent=root)
        root.destroy()
    except Exception:
        # This fallback is useful for headless and command-line execution.
        pass


def choose_ranking_file() -> Path:
    """Open the platform file chooser only when an input path was not supplied."""
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        show_fullscreen_message(
            root,
            "Import ranking data",
            "You are about to import the ranking data for the training program "
            "matching software. This requires specific formatting.\n\n"
            "1) The spreadsheet must be saved as either a .csv or .xlsx file. If it "
            "uses another format, go into your spreadsheet software and use 'Save as' "
            "to change the file type to .csv or .xlsx.\n\n"
            "2) The first row of the spreadsheet should be headings; the rest of the "
            "rows must each correspond to a different trainee.\n\n"
            "3) Certain columns must contain specific information:\n"
            "    Column C must contain the trainee's name.\n"
            "    Column D must declare whether they are 'full-time' or 'part-time' and "
            "say this text explicitly.\n"
            "    Column F contains their first choice site, G their second, and so on "
            "until their last choice.\n\n"
            "Any deviation from the above requirements will cause the program not to "
            "run properly.",
            use_parent=True,
            require_ok=True,
        )
        root.withdraw()
        root.attributes("-topmost", True)
        selected = filedialog.askopenfilename(
            title="Select doctor ranking spreadsheet",
            filetypes=[("Spreadsheet files", "*.csv *.xlsx"), ("CSV files", "*.csv"),
                       ("Excel files", "*.xlsx")],
        )
        root.destroy()
    except Exception as error:
        raise MatchingError(
            "Could not open the file browser because this Python installation has no usable "
            f"Tkinter GUI support ({error}). Install Tkinter for this Python, or supply "
            "the spreadsheet with --rankings PATH instead."
        ) from error
    if not selected:
        raise MatchingError("No ranking spreadsheet was selected.")
    return Path(selected)


def choose_output_file() -> Path:
    """Ask where to save the results when the app is launched graphically."""
    try:
        import tkinter as tk
        from tkinter import filedialog, ttk

        root = tk.Tk()
        root.title("Output format")
        root.resizable(False, False)
        format_name = tk.StringVar(value="CSV (.csv)")
        selected_format: list[str] = []
        panel = ttk.Frame(root, padding=18)
        panel.grid(sticky="nsew")
        ttk.Label(panel, text="Save assignments as:").grid(row=0, column=0, padx=(0, 8), sticky="w")
        format_choice = ttk.Combobox(panel, textvariable=format_name, state="readonly", width=18,
                                     values=("CSV (.csv)", "Excel (.xlsx)"))
        format_choice.grid(row=0, column=1, sticky="ew")
        panel.columnconfigure(1, weight=1)
        completed = tk.BooleanVar(master=root, value=False)

        def continue_to_save() -> None:
            selected_format.append(format_name.get())
            completed.set(True)

        def cancel() -> None:
            completed.set(True)

        actions = ttk.Frame(panel)
        actions.grid(row=1, column=0, columnspan=2, pady=(16, 0), sticky="e")
        ttk.Button(actions, text="Cancel", command=cancel).grid(row=0, column=0, padx=(0, 6))
        ttk.Button(actions, text="Choose save location", command=continue_to_save).grid(row=0, column=1)
        root.protocol("WM_DELETE_WINDOW", cancel)
        root.wait_variable(completed)
        for child in root.winfo_children():
            child.destroy()

        if not selected_format:
            selected = ""
        else:
            excel_output = selected_format[0] == "Excel (.xlsx)"
            extension = ".xlsx" if excel_output else ".csv"
            file_type = ("Excel files", "*.xlsx") if excel_output else ("CSV files", "*.csv")
            root.withdraw()
            root.attributes("-topmost", True)
            selected = filedialog.asksaveasfilename(
                title="Save hospital assignments",
                defaultextension=extension,
                filetypes=[file_type],
            )
        root.destroy()
    except Exception as error:
        raise MatchingError(
            "Could not open the save-file browser because this Python installation has "
            f"no usable Tkinter GUI support ({error}). Supply an output path instead."
        ) from error
    if not selected:
        raise MatchingError("No output file was selected.")
    return Path(selected)


def choose_director_settings(doctors: list[Doctor]) -> tuple[dict[str, Hospital], dict[str, str]]:
    """Collect hospital capacity and locked assignments in a small native form."""
    try:
        import tkinter as tk
        from tkinter import messagebox, ttk
    except ImportError as error:
        raise MatchingError(
            "The director form requires Tkinter, which is unavailable in this Python "
            f"installation ({error}). Use --config PATH instead."
        ) from error

    try:
        root = tk.Tk()
    except Exception as error:
        raise MatchingError(
            f"Could not open the director form ({error}). Use --config PATH instead."
        ) from error
    root.title("Doctor–hospital matcher: director settings")
    root.minsize(720, 520)
    result: tuple[dict[str, Hospital], dict[str, str]] | None = None
    # Preserve the first-seen order from the spreadsheet. Capacity remains zero until
    # the director explicitly enters it, preventing an accidental allocation.
    hospitals: dict[str, Hospital] = {}
    for doctor in doctors:
        for name in doctor.rankings:
            hospitals.setdefault(name, Hospital(name, 0, False))
    locks: dict[str, str] = {}

    frame = ttk.Frame(root, padding=12)
    frame.grid(sticky="nsew")
    root.columnconfigure(0, weight=1)
    root.rowconfigure(0, weight=1)
    frame.columnconfigure(0, weight=1)
    frame.rowconfigure(1, weight=1)
    frame.rowconfigure(4, weight=1)

    heading = ttk.Frame(frame)
    heading.grid(row=0, column=0, sticky="ew")
    heading.columnconfigure(0, weight=1)
    ttk.Label(heading, text="Hospitals", font=("TkDefaultFont", 11, "bold")).grid(row=0, column=0, sticky="w")
    hospital_tree = ttk.Treeview(frame, columns=("places", "jobshares"), show="tree headings", height=7)
    hospital_tree.heading("#0", text="Hospital")
    hospital_tree.heading("places", text="Full-time places")
    hospital_tree.heading("jobshares", text="Allows jobshares")
    hospital_tree.column("#0", width=290)
    hospital_tree.column("places", width=130, anchor="center")
    hospital_tree.column("jobshares", width=130, anchor="center")
    hospital_tree.grid(row=1, column=0, sticky="nsew")
    hospital_entry_row = ttk.Frame(frame)
    hospital_entry_row.grid(row=2, column=0, pady=(6, 14), sticky="ew")
    ttk.Label(hospital_entry_row, text="Name").grid(row=0, column=0)
    hospital_name = ttk.Entry(hospital_entry_row, width=28)
    hospital_name.grid(row=0, column=1, padx=(4, 10))
    ttk.Label(hospital_entry_row, text="Places").grid(row=0, column=2)
    hospital_places = ttk.Entry(hospital_entry_row, width=7)
    hospital_places.grid(row=0, column=3, padx=(4, 10))
    jobshares = tk.BooleanVar(value=False)
    ttk.Checkbutton(hospital_entry_row, text="Allows jobshares", variable=jobshares).grid(row=0, column=4)

    ttk.Label(frame, text="Pre-allocations (optional)", font=("TkDefaultFont", 11, "bold")).grid(row=3, column=0, sticky="w")
    lock_tree = ttk.Treeview(frame, columns=("hospital",), show="tree headings", height=7)
    lock_tree.heading("#0", text="Doctor")
    lock_tree.heading("hospital", text="Allocated hospital")
    lock_tree.column("#0", width=330)
    lock_tree.column("hospital", width=280)
    lock_tree.grid(row=4, column=0, sticky="nsew")
    lock_entry_row = ttk.Frame(frame)
    lock_entry_row.grid(row=5, column=0, pady=(6, 12), sticky="ew")
    ttk.Label(lock_entry_row, text="Doctor").grid(row=0, column=0)
    doctor_choice = ttk.Combobox(lock_entry_row, values=[doctor.name for doctor in doctors], state="readonly", width=28)
    doctor_choice.grid(row=0, column=1, padx=(4, 10))
    ttk.Label(lock_entry_row, text="Hospital").grid(row=0, column=2)
    lock_hospital_choice = ttk.Combobox(lock_entry_row, state="readonly", width=28)
    lock_hospital_choice.grid(row=0, column=3, padx=(4, 10))

    def refresh() -> None:
        for item in hospital_tree.get_children():
            hospital_tree.delete(item)
        for name, hospital in hospitals.items():
            hospital_tree.insert("", "end", iid=name, text=name,
                                 values=(hospital.places, "Yes" if hospital.allows_jobshares else "No"))
        for item in lock_tree.get_children():
            lock_tree.delete(item)
        for doctor, hospital in locks.items():
            lock_tree.insert("", "end", iid=doctor, text=doctor, values=(hospital,))
        lock_hospital_choice["values"] = list(hospitals)

    def load_selected_hospital(_event: object | None = None) -> None:
        selected = hospital_tree.selection()
        if not selected:
            return
        hospital = hospitals[selected[0]]
        hospital_name.delete(0, "end")
        hospital_name.insert(0, hospital.name)
        hospital_places.delete(0, "end")
        hospital_places.insert(0, str(hospital.places))
        jobshares.set(hospital.allows_jobshares)

    def add_hospital() -> None:
        name = hospital_name.get().strip()
        try:
            places = int(hospital_places.get().strip())
            if not name or places < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Invalid hospital", "Enter a hospital name and a non-negative whole number of places.", parent=root)
            return
        selected = hospital_tree.selection()
        previous_name = selected[0] if selected else None
        if previous_name and previous_name != name:
            if name in hospitals:
                messagebox.showerror("Duplicate hospital", "A hospital with that name already exists.", parent=root)
                return
            # A name imported from the spreadsheet is also present in doctors'
            # rankings, so renaming it must update those references as well.
            del hospitals[previous_name]
            for index, doctor in enumerate(doctors):
                updated_rankings = tuple(
                    dict.fromkeys(name if ranked == previous_name else ranked for ranked in doctor.rankings)
                )
                doctors[index] = Doctor(doctor.name, doctor.employment, updated_rankings, doctor.row_number)
            for doctor_name, locked_hospital in list(locks.items()):
                if locked_hospital == previous_name:
                    locks[doctor_name] = name
        hospitals[name] = Hospital(name, places, jobshares.get())
        hospital_name.delete(0, "end")
        hospital_places.delete(0, "end")
        refresh()

    def remove_hospital() -> None:
        selected = hospital_tree.selection()
        if not selected:
            return
        name = selected[0]
        del hospitals[name]
        for doctor in [doctor for doctor, hospital in locks.items() if hospital == name]:
            del locks[doctor]
        refresh()

    def add_lock() -> None:
        doctor, hospital = doctor_choice.get(), lock_hospital_choice.get()
        if not doctor or not hospital:
            messagebox.showerror("Incomplete pre-allocation", "Choose both a doctor and a hospital.", parent=root)
            return
        locks[doctor] = hospital
        refresh()

    def remove_lock() -> None:
        selected = lock_tree.selection()
        if selected:
            del locks[selected[0]]
            refresh()

    ttk.Button(hospital_entry_row, text="Add / update", command=add_hospital).grid(row=0, column=5, padx=(10, 4))
    ttk.Button(hospital_entry_row, text="Remove selected", command=remove_hospital).grid(row=0, column=6)
    ttk.Button(lock_entry_row, text="Allocate", command=add_lock).grid(row=0, column=4, padx=(10, 4))
    ttk.Button(lock_entry_row, text="Remove selected", command=remove_lock).grid(row=0, column=5)

    actions = ttk.Frame(frame)
    actions.grid(row=6, column=0, sticky="e")

    def show_help() -> None:
        show_fullscreen_message(
            root,
            "How to use the matcher",
            "1. The spreadsheet you selected contains the doctors and their ranked "
            "hospital choices. Doctor names and rankings cannot be edited here.\n\n"
            "2. Under Hospitals, names found in the doctors' rankings are imported "
            "automatically with zero places. Select a hospital to load its editable "
            "details below, set its number of full-time places and jobshare rule, then "
            "click Add / update. You can also enter a new name to add another hospital. "
            "You can also add hospitals that were not listed by any doctors. To deselect "
            "a hospital that you have selected, just press Add / update having made no "
            "changes. Setting 'Allows Jobshares' to 'Yes' means that the hospital allows "
            "for two part-time doctors to share one full-time job.\n\n"
            "3. Pre-allocations allow the TPD to ensure a doctor is assigned to a "
            "particular hospital, e.g. for reasons they have discussed privately. This "
            "section can only be filled out after all the hospital information was "
            "provided above.\n\n"
            "4. Select an entry in either table and click Remove selected to delete it.\n\n"
            "5. Click Run matching. You will then choose where to save the output CSV "
            "or XLSX file. "
            "Each column in that CSV is a hospital and lists its assigned doctors.\n\n"
            "If a doctor cannot be assigned to one of their ranked hospitals, they are "
            "listed in the UNALLOCATED column of the output.",
        )

    def save() -> None:
        nonlocal result
        if not hospitals:
            messagebox.showerror("Hospitals required", "Add at least one hospital.", parent=root)
            return
        try:
            # Validate before closing the form so the director can correct capacity,
            # ranking, or pre-allocation issues without re-entering everything.
            _assignments, unmatched = match(doctors, hospitals, locks)
        except MatchingError as error:
            messagebox.showerror("Cannot run matching", str(error), parent=root)
            return
        if unmatched:
            save_anyway = messagebox.askyesno(
                "Unallocated doctors",
                "The following doctor(s) could not be allocated to any ranked hospital:\n\n"
                + ", ".join(unmatched)
                + "\n\nThey will be placed in an UNALLOCATED column in the output. "
                "Do you want to save the matching anyway?",
                parent=root,
            )
            if not save_anyway:
                return
        result = hospitals.copy(), locks.copy()
        root.destroy()
    ttk.Button(actions, text="Cancel", command=root.destroy).grid(row=0, column=0, padx=4)
    ttk.Button(actions, text="Help", command=show_help).grid(row=0, column=1, padx=4)
    ttk.Button(actions, text="Run matching", command=save).grid(row=0, column=2)
    hospital_tree.bind("<<TreeviewSelect>>", load_selected_hospital)
    refresh()
    root.protocol("WM_DELETE_WINDOW", root.destroy)
    root.mainloop()
    if result is None:
        raise MatchingError("Director setup was cancelled.")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Match doctors to hospitals using their ranked choices.")
    parser.add_argument("output", nargs="?", type=Path,
                        help="output assignments CSV; choose it in the GUI when omitted")
    parser.add_argument("--rankings", type=Path, metavar="PATH",
                        help="CSV or XLSX ranking spreadsheet; skips the file browser")
    parser.add_argument("--config", type=Path, metavar="PATH",
                        help="director configuration JSON; skips the director form")
    parser.add_argument("--no-header", action="store_true", help="treat the first row as a doctor rather than headings")
    args = parser.parse_args()
    try:
        rankings = args.rankings or choose_ranking_file()
        doctors = read_doctors(rankings, skip_header=not args.no_header)
        hospitals, preallocations = (read_configuration(args.config) if args.config
                                     else choose_director_settings(doctors))
        assignments, unmatched = match(doctors, hospitals, preallocations)
        output = args.output or choose_output_file()
        write_output(output, hospitals, assignments, unmatched)
    except MatchingError as error:
        message = str(error)
        print(f"Error: {message}", file=sys.stderr)
        show_error(message)
        return 2
    print(f"Wrote {output}.")
    if unmatched:
        print("Unmatched doctors: " + ", ".join(unmatched), file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
